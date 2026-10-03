"""Referentiel des titres (Maroclear) : chargement depuis le CSV livre ou un fichier Excel / CSV de l utilisateur."""
from __future__ import annotations
import csv, io, datetime as dt
from .bdt import Titre


def _date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _num(v, defaut=0.0):
    if v is None or v == "":
        return defaut
    try:
        return float(str(v).replace(",", ".").replace("%", "").replace(" ", ""))
    except ValueError:
        return defaut


def seulement_bdt(ref: dict) -> dict:
    return {k: t for k, t in ref.items() if "BDT" in (t.libelle or "").upper()}


def charger_csv(chemin_ou_texte, est_texte=False) -> dict:
    """CSV (;) avec en-tete : code;isin;libelle;emetteur;categorie;nominal;taux_facial;jouissance;echeance;periodicite;amortissement;spread_bps"""
    f = io.StringIO(chemin_ou_texte) if est_texte else open(chemin_ou_texte, encoding="utf-8")
    with f:
        rd = csv.DictReader(f, delimiter=";")
        out = {}
        for r in rd:
            try:
                code = str(r["code"]).strip()
                tf = _num(r.get("taux_facial"))
                if tf > 1:
                    tf = tf / 100
                t = Titre(code=code, libelle=r.get("libelle", ""), nominal=_num(r.get("nominal"), 100000), taux_facial=tf,
                          jouissance=_date(r.get("jouissance")), echeance=_date(r.get("echeance")), spread_bps=0.0,
                          isin=r.get("isin", "") or "")
                if t.echeance and t.jouissance:
                    out[code] = t
            except Exception:
                continue
    return out


def charger_excel(fichier) -> dict:
    """Feuille Referentiel du classeur (B:M a partir de la ligne 7), ou tout tableau aux memes colonnes."""
    import openpyxl
    wb = openpyxl.load_workbook(fichier, data_only=True)
    ws = wb["Referentiel"] if "Referentiel" in wb.sheetnames else wb.worksheets[0]
    out = {}
    for row in ws.iter_rows(min_row=7, values_only=True):
        row = list(row) + [None] * 14
        code = row[1]
        if code is None:
            continue
        try:
            tf = _num(row[7])
            if tf > 1:
                tf = tf / 100
            t = Titre(code=str(code).strip(), libelle=str(row[3] or ""), nominal=_num(row[6], 100000), taux_facial=tf,
                      jouissance=_date(row[8]), echeance=_date(row[9]), spread_bps=0.0, isin=str(row[2] or ""))
            if t.echeance and t.jouissance:
                out[t.code] = t
        except Exception:
            continue
    return out
